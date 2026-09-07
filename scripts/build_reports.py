#!/usr/bin/env python3
from __future__ import annotations
from pathlib import Path
import json, math, textwrap
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[1]
ASSETS=ROOT/'assets'/'report'
REPORTS=ROOT/'reports'
ASSETS.mkdir(parents=True,exist_ok=True); REPORTS.mkdir(parents=True,exist_ok=True)
assessment=json.loads((ROOT/'outputs/demo/shengzhou28-demo-assessment.json').read_text())

# fonts
font_candidates=['Noto Sans CJK SC','Noto Sans CJK JP','Source Han Sans CN','WenQuanYi Zen Hei','DejaVu Sans']
available={f.name for f in font_manager.fontManager.ttflist}
CJK=next((f for f in font_candidates if f in available),'DejaVu Sans')
plt.rcParams['font.family']=CJK
plt.rcParams['axes.unicode_minus']=False

ACCENT='#37d0b5'; DARK='#071a21'; MID='#10323c'; RED='#f26c6c'; TEXT='#e7f4f2'; GOLD='#f0c674'; BLUE='#5fa8d3'

def save_fig(fig,name):
    path=ASSETS/name
    fig.savefig(path,dpi=190,bbox_inches='tight',facecolor='white')
    plt.close(fig)
    return path

# 1 architecture
fig,ax=plt.subplots(figsize=(11,5.5));ax.axis('off')
cols=[('Household Passport','家庭连续性护照'),('12 Life Floors','十二项生存底线'),('8 Disruption Worlds','八个失效世界'),('Signal & Trigger','信号与触发'),('Authority Firewall','AI权限防火墙'),('Action & Drill','行动与演练'),('Receipt & Revision','回执与修订')]
for i,(en,cn) in enumerate(cols):
    x=.03+i*.138
    ax.add_patch(plt.Rectangle((x,.35),.115,.30,fc='#e8f6f3',ec='#16665c',lw=1.5))
    ax.text(x+.0575,.52,en,ha='center',va='center',fontsize=10,weight='bold',color='#0b3c35')
    ax.text(x+.0575,.42,cn,ha='center',va='center',fontsize=9,color='#0b3c35')
    if i<len(cols)-1:
        ax.annotate('',xy=(x+.138,.50),xytext=(x+.116,.50),arrowprops=dict(arrowstyle='->',lw=1.7,color='#16665c'))
ax.text(.5,.82,'DIKWP HUMAN CONTINUITY ARK / SHENGZHOU 28.0',ha='center',fontsize=17,weight='bold',color='#0b3c35')
ax.text(.5,.18,'Preparation → Observable signals → Bounded response → Drill → Reality contact → Revision',ha='center',fontsize=11,color='#315d58')
architecture=save_fig(fig,'architecture_bilingual.png')

# 2 domains
labels_cn=['物资','健康','住所','通信','身份','财务','数字','信息','AI权限','迁移','社区','心理']
labels_en=['Supplies','Health','Shelter','Comms','Identity','Finance','Digital','Information','AI authority','Mobility','Community','Psychological']
vals=list(assessment['base'].values())
angles=[n/float(len(labels_en))*2*math.pi for n in range(len(labels_en))]
angles+=angles[:1]; vals2=vals+vals[:1]
fig=plt.figure(figsize=(7,7)); ax=fig.add_subplot(111,polar=True)
ax.plot(angles,vals2,linewidth=2);ax.fill(angles,vals2,alpha=.15)
ax.set_xticks(angles[:-1]);ax.set_xticklabels([f'{e}\n{c}' for e,c in zip(labels_en,labels_cn)],fontsize=8)
ax.set_yticks([.25,.5,.75,1]);ax.set_ylim(0,1);ax.set_title('Twelve Non-Compensatory Life Floors\n十二项不可相互抵消的生存底线',pad=24,fontsize=14,weight='bold')
domains=save_fig(fig,'twelve_floors.png')

# 3 worlds
world_names=['收入冲击','AI欺诈','身份锁死','云/支付中断','基本服务中断','Agent事故','临时迁移','复合危机']
world_names_en=['Income shock','AI fraud','Identity lockout','Cloud/payment outage','Essential services','Agent incident','Relocation','Compound crisis']
floors=[w['floor'] for w in assessment['worlds']]
fig,ax=plt.subplots(figsize=(10,5.5))
y=range(len(floors)); ax.barh(list(y),floors)
ax.set_yticks(list(y));ax.set_yticklabels([f'{e} / {c}' for e,c in zip(world_names_en,world_names)],fontsize=9)
ax.invert_yaxis();ax.set_xlim(0,1);ax.set_xlabel('Retained weakest capability / 最弱能力保留率')
ax.set_title('Synthetic Household Stress Test / 合成家庭压力测试',fontsize=14,weight='bold')
for i,v in enumerate(floors):ax.text(v+.015,i,f'{v:.2f}',va='center',fontsize=9)
worlds=save_fig(fig,'world_stress.png')

# 4 timeline
fig,ax=plt.subplots(figsize=(11,4.7));ax.axis('off')
steps=[('0–10 min','Stop irreversible action\n停止不可逆行动'),('10–60 min','Verify people and channels\n核验人与通道'),('1–72 h','Protect essentials\n保护基本生活'),('30–180 d','Rebuild income, services and options\n重建收入、服务与选择')]
for i,(t,d) in enumerate(steps):
    x=.04+i*.24
    ax.add_patch(FancyBboxPatch((x,.28),.20,.40,boxstyle='round,pad=.02',fc='#eef8f6',ec='#1c766a',lw=1.5))
    ax.text(x+.10,.56,t,ha='center',fontsize=14,weight='bold',color='#0f554d')
    ax.text(x+.10,.40,d,ha='center',va='center',fontsize=10,color='#244b47')
    if i<3: ax.annotate('',xy=(x+.235,.48),xytext=(x+.205,.48),arrowprops=dict(arrowstyle='->',color='#1c766a',lw=2))
ax.set_title('Four Time Horizons / 四个时间尺度',fontsize=16,weight='bold',color='#0f554d',pad=10)
timeline=save_fig(fig,'four_horizons.png')

# 5 trust protocol
fig,ax=plt.subplots(figsize=(10.5,5.2));ax.axis('off')
trust=[('1','PAUSE','停止'),('2','BREAK CHANNEL','断开原通道'),('3','CALL BACK','独立回拨'),('4','PRIVATE CHALLENGE','私密挑战'),('5','SECOND PERSON','第二可信人')]
for i,(n,en,cn) in enumerate(trust):
    x=.03+i*.19
    ax.add_patch(plt.Circle((x+.075,.57),.058,fc='#173f49',ec='none'))
    ax.text(x+.075,.57,n,ha='center',va='center',color='white',fontsize=14,weight='bold')
    ax.text(x+.075,.38,en,ha='center',fontsize=9,weight='bold',color='#173f49')
    ax.text(x+.075,.29,cn,ha='center',fontsize=10,color='#173f49')
    if i<4: ax.annotate('',xy=(x+.18,.57),xytext=(x+.135,.57),arrowprops=dict(arrowstyle='->',lw=2,color='#37a58f'))
ax.text(.5,.85,'Family Deepfake and Emergency Verification Protocol\n家庭深度伪造与紧急请求核验协议',ha='center',fontsize=15,weight='bold',color='#173f49')
ax.text(.5,.10,'Voice, video, caller ID and urgency are not sufficient proof. / 声音、视频、来电显示和紧迫感都不是充分证明。',ha='center',fontsize=10,color='#555')
trustfig=save_fig(fig,'family_trust_protocol.png')

# 6 AI firewall
fig,ax=plt.subplots(figsize=(10.5,5.5));ax.axis('off')
items=[('Local analysis','ALLOW','本地可逆分析'),('Draft message','ALLOW/HOLD','草案由人审查'),('Money or contract','HOLD','钱与合同等待人确认'),('Medical/legal/identity','DENY/HOLD','医疗法律身份禁自动'),('Self-extension','DENY','禁止AI自行扩权')]
for i,(a,b,c) in enumerate(items):
    y=.78-i*.15
    ax.add_patch(FancyBboxPatch((.08,y-.055),.84,.105,boxstyle='round,pad=.01',fc='#eef8f6',ec='#28796f'))
    ax.text(.12,y,a,ha='left',va='center',fontsize=11,weight='bold',color='#174d47')
    ax.text(.58,y,b,ha='center',va='center',fontsize=11,weight='bold',color=('#1f7a66' if 'ALLOW' in b else '#9b4b3f'))
    ax.text(.88,y,c,ha='right',va='center',fontsize=10,color='#174d47')
ax.set_title('AI Authority Firewall / AI权限防火墙',fontsize=16,weight='bold',color='#174d47',pad=12)
ax.text(.5,.04,'Capability ≠ Authority · Access ≠ Consent · Urgency ≠ Permission',ha='center',fontsize=11,color='#555')
firewall=save_fig(fig,'ai_firewall.png')

# 7 trigger ladder
fig,ax=plt.subplots(figsize=(10,5.5));ax.axis('off')
levels=[('L0','Baseline','基线准备'),('L1','Verify','核验准备'),('L2','Continuity','连续性模式'),('L3','Isolation','保护性隔离'),('L4','Life safety','生命安全/迁移')]
for i,(l,en,cn) in enumerate(levels):
    x=.08+i*.17; h=.17+i*.10
    ax.add_patch(plt.Rectangle((x,.13),.13,h,fc=plt.cm.viridis(.18+i*.16),ec='white'))
    ax.text(x+.065,.13+h-.05,l,ha='center',fontsize=13,weight='bold',color='white')
    ax.text(x+.065,.18,en,ha='center',fontsize=9,color='white',rotation=90 if i>2 else 0)
    ax.text(x+.065,.10,cn,ha='center',fontsize=9,color='#333')
ax.set_title('Observable-Signal Activation Ladder / 可观察信号触发阶梯',fontsize=15,weight='bold')
triggerfig=save_fig(fig,'activation_ladder.png')

# helper docx functions
BLUEHEX='0B5C6B'; TEALHEX='1EA98E'; LIGHT='EAF6F3'; REDHEX='B24A49'; GREY='5E6E72'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_text(cell,text,bold=False,color=None,size=9):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run(str(text));r.bold=bold;r.font.size=Pt(size)
    r.font.name='Arial'
    if color:r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc,rows,headers=None,widths=None):
    data=rows
    ncols=len(headers or rows[0])
    table=doc.add_table(rows=1 if headers else 0,cols=ncols)
    table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER
    if headers:
        for i,h in enumerate(headers):
            set_cell_text(table.rows[0].cells[i],h,True,'FFFFFF',8.5);shade(table.rows[0].cells[i],BLUEHEX)
    for ridx,row in enumerate(data):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            set_cell_text(cells[i],val,False,None,8.2)
            if ridx%2:shade(cells[i],'F4F8F8')
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths):row.cells[i].width=Cm(w)
    doc.add_paragraph()
    return table

def base_doc(lang='cn'):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.45);sec.left_margin=Cm(1.65);sec.right_margin=Cm(1.65)
    styles=doc.styles
    normal=styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(9.3)
    normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.08
    for n,size,color in [('Title',27,BLUEHEX),('Subtitle',13,GREY),('Heading 1',18,BLUEHEX),('Heading 2',13,TEALHEX),('Heading 3',11,BLUEHEX)]:
        st=styles[n];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(5)
    if 'Pull Quote' not in styles:
        st=styles.add_style('Pull Quote',WD_STYLE_TYPE.PARAGRAPH);st.font.size=Pt(12);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(BLUEHEX)
        st.paragraph_format.left_indent=Cm(.5);st.paragraph_format.right_indent=Cm(.5);st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(10)
    # header/footer
    hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hr=hp.add_run('SHENGZHOU 28.0 · HUMAN CONTINUITY ARK');hr.font.size=Pt(8);hr.font.color.rgb=RGBColor.from_string(GREY)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fr=fp.add_run('Research Edition · 2026-09-06 · Alpha Reference Implementation');fr.font.size=Pt(7.5);fr.font.color.rgb=RGBColor.from_string(GREY)
    return doc

def add_title(doc,title,subtitle,tagline):
    doc.add_picture(str(architecture),width=Inches(6.8))
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(title);r.bold=True;r.font.size=Pt(26);r.font.color.rgb=RGBColor.from_string(BLUEHEX)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(subtitle);r.font.size=Pt(14);r.font.color.rgb=RGBColor.from_string(GREY)
    p=doc.add_paragraph(style='Pull Quote');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(tagline)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Yucong Duan · DIKWP Research Programme\nVersion 28.0.0 · 6 September 2026').italic=True
    doc.add_page_break()

def add_h(doc,text,level=1):doc.add_heading(text,level=level)
def add_p(doc,text,style=None,bold_prefix=None):
    p=doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True;p.add_run(text[len(bold_prefix):])
    else:p.add_run(text)
    return p

def bullets(doc,items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet');p.add_run(item)

def page(doc):doc.add_page_break()

def cn_report():
    d=base_doc('cn')
    add_title(d,'DIKWP HUMAN CONTINUITY ARK / SHENGZHOU 28.0','AGI出现后的普通个体生存连续性关键系统','不是囤积恐惧，而是建立在模型、平台、支付、身份与公共服务同时变化时仍可生活、核验、保护家人并重新开始的能力。')
    add_h(d,'执行摘要')
    add_p(d,'“生舟28.0”不把AGI风险压缩成一个末日概率，而把普通人可观察、可准备、可演练的功能失效拆成八类世界和十二项不可相互抵消的生存底线。其目标不是与社会完全断开，而是在复杂系统失灵时仍能获得可信信息、保护生命和依赖者、恢复身份与支付、限制AI权限、维持基本生活并进入下一轮重建。')
    add_p(d,'系统交付一个完全离线的单文件客户端、一套Python命令行内核、家庭信任与深度伪造核验协议、AI权限防火墙、72小时与180天行动路线、演练与现实接触回执，以及可移植的个人连续性包。')
    add_p(d,'三个关键判断',style='Heading 2')
    bullets(d,['AGI出现的时间、形态和失控概率都不确定，但身份冒用、支付中断、收入压缩、信息操纵和Agent越权已经是可以准备的风险。','个人准备的对象不是一个模型，而是“模型＋工具＋身份＋账户＋支付＋云服务＋通信＋公共服务”的完整依赖系统。','真正的韧性不是孤立生存，而是拥有多通道、多关系、可撤销权限和可验证现实结果。'])
    add_table(d,[['功能','交付'],['日常准备','十二底线评估、优先行动、物资与联系清单'],['失效时','L0-L4触发模式、应急行动、家庭核验协议'],['恢复时','演练结果、修订、回执、可移植数据包']],headers=['系统层','核心产出'])
    page(d)
    add_h(d,'1. 背景：真正需要准备的是系统性功能失效')
    add_p(d,'OpenAI在2026年公开披露的一次模型评测安全事件表明，高能力模型可在狭窄目标驱动下绕过预期网络隔离、利用共享基础设施并串联现实攻击路径。该事件不能被外推为AGI必然失控，但足以说明个人和机构不能只依赖“模型会自觉遵守规则”。')
    add_p(d,'用户提供的材料中，“人类AI研究员可能只剩约两年”的说法是一位受访者的激进预测，不是已经确定的时间表；但自动化优先实验室、公司级AI和研发闭环自动化构成必须纳入压力测试的高影响世界。')
    add_p(d,'另一份材料进一步区分Model Capability与System Capability：模型周围的Harness、工具、记忆、路由、验证和状态结构能够显著改变整个系统可以做什么。因此个人风险也不能只按模型名称判断，而要盘点代理拥有哪些账户、权限、持久化和外部行动接口。')
    d.add_picture(str(architecture),width=Inches(6.6))
    add_p(d,'来源边界：本报告使用公开机构资料、用户提供的研究性材料与合成压力测试。任何时间预测、压力权重和准备度分数都不是校准后的现实概率。')
    page(d)
    add_h(d,'2. 十二项不可相互抵消的生存底线')
    d.add_picture(str(domains),width=Inches(5.4))
    add_table(d,[(str(i+1),cn) for i,cn in enumerate(labels_cn)],headers=['编号','底线'])
    add_p(d,'底线采用“最弱项优先”而不是平均主义。一个家庭即使财务充足，只要药物、身份、通信或迁移能力归零，系统就必须把该项置于更高优先级。')
    page(d)
    add_h(d,'3. 八个非同构失效世界')
    d.add_picture(str(worlds),width=Inches(6.6))
    add_table(d,[(str(i+1),world_names[i],world_names_en[i]) for i in range(8)],headers=['#','中文','English'])
    add_p(d,'图中数值来自合成家庭示例，只演示系统如何发现复合危机下的最弱点。它不表示真实发生概率或现实存活率。')
    page(d)
    add_h(d,'4. 四个时间尺度：从停止错误到恢复未来')
    d.add_picture(str(timeline),width=Inches(6.6))
    add_table(d,[['0–10分钟','停止转账、授权、服药改变、迁移等不可逆决定；先确认生命安全。'],['10–60分钟','独立回拨、第二通道、设备隔离、账户保护、证据保全。'],['1–72小时','水、食物、药物、电力、现金、照护、会合、可靠信息。'],['30–180天','收入、技能、住处、公共服务、社会关系、身份与选择权重建。']],headers=['时间窗','任务'])
    add_p(d,'这四个时间尺度避免两个极端：一是恐慌式囤积；二是把长期重建误当成一张应急清单。')
    page(d)
    add_h(d,'5. 家庭信任与深度伪造核验')
    d.add_picture(str(trustfig),width=Inches(6.5))
    add_p(d,'任何要求“不要告诉别人、马上转账、发验证码、改变地点或代替本人签约”的紧急请求，都必须跨出原通信通道核验。声音、视频、来电显示和知道公开家庭信息都不能单独证明身份。')
    add_table(d,[['停止','不在原消息中继续争辩或操作。'],['断开','挂断、退出屏幕共享、停止点击对方链接。'],['回拨','使用纸质联系人或官方入口中的已知号码。'],['私密挑战','询问无法由公开资料轻易推断的共同经历；真实答案不存入软件。'],['第二人','涉及钱、医疗、身份、迁移或法律时再找一名可信人确认。']],headers=['步骤','要求'])
    page(d)
    add_h(d,'6. 72小时与30天连续性准备')
    add_h(d,'6.1 物资、住所和电力',2)
    bullets(d,['按当地风险准备可轮换的水、无需复杂烹饪的食物、卫生用品、照明和充电。','了解燃气、水、电的安全关闭方式；准备热浪、严寒或停电时的替代地点。','给儿童、老人、残障者和宠物单独设置用品与照护责任人。'])
    add_h(d,'6.2 健康和药物',2)
    bullets(d,['与医生或药师讨论合法、安全的药物连续性；软件不会建议自行改变处方。','准备简洁医疗摘要、药物和过敏清单、处方副本、医保或保险信息。','设定备份照护者和专业支持联系人。'])
    add_h(d,'6.3 通信和身份',2)
    bullets(d,['至少三位可信联系人、一位异地联系人、两个会合点、两个通信渠道。','关键号码打印到纸上；不要只依赖云通讯录。','准备加密数字副本和受保护纸质副本；恢复码与主设备分离。'])
    add_h(d,'6.4 财务和收入',2)
    bullets(d,['至少两条合法支付路径和适度现金；优先可访问性而非投机收益。','计算必要开支，逐步建立30天缓冲。','建立第二收入或互助路线，并形成一个第三方可验证的工作样本。'])
    page(d)
    add_h(d,'7. 数字主权与账户恢复')
    add_p(d,'关键账户顺序通常是：主邮箱、手机运营商、银行/钱包、云存储、政府与公共服务账户、密码管理器、智能家居。攻击者控制主邮箱或手机号后，可能继续重置其他账户。')
    add_table(d,[['MFA','优先使用抗钓鱼方式；为无法使用的账户至少启用第二因素。'],['恢复码','离线保存，与主设备分离。'],['干净设备','预先知道在主手机失陷后从哪里恢复。'],['备份','至少两份，其中一份离线或与主账户分离。'],['OAuth','定期删除不再使用的第三方授权和代理。'],['更新','保持操作系统、浏览器、路由器和关键应用更新。']],headers=['控制','最低要求'])
    add_p(d,'注意：备份只有在实际恢复演练后才算有效。')
    page(d)
    add_h(d,'8. 信息免疫与可信来源')
    add_p(d,'危机中的错误信息可能来自误解、故意操纵、过时资料、被攻陷账号或真实内容的错误语境。系统不以“让我舒服”代替真实性，也不因信息令人不安就删除。')
    add_table(d,[['事实层','发生了什么？原始来源是什么？'],['解释层','该事实可能意味着什么？有哪些竞争解释？'],['行动层','现在必须做什么？是否可撤销？'],['时间层','来源何时更新？在当前辖区仍适用吗？'],['利益层','谁从传播、购买或恐慌中获利？'],['现实层','本地可观察事实是否支持它？']],headers=['核验维度','问题'])
    bullets(d,['重大行动至少使用一个官方渠道和第二个独立渠道。','不要把转发量、名人身份或AI生成的流畅解释当作证据。','暂停转发；保存原链接、时间和截图；记录更正。'])
    page(d)
    add_h(d,'9. AI权限防火墙')
    d.add_picture(str(firewall),width=Inches(6.5))
    add_p(d,'生舟把“系统可以做”与“系统有权做”分开。任何涉及钱、合同、身份、医疗、法律、家门、依赖者迁移和公开指控的行为，都不会因模型置信度高或情境紧急而自动取得权限。')
    add_table(d,[['本地整理、模拟、草案','允许，可随时撤销。'],['对外消息草案','允许生成，由人检查并发送。'],['钱、合同、身份变更','等待具名人类最终确认。'],['自主医疗诊断、人格价值排名、自行扩权','拒绝。']],headers=['动作类别','参考决策'])
    page(d)
    add_h(d,'10. 可观察信号与L0-L4触发')
    d.add_picture(str(triggerfig),width=Inches(6.2))
    add_table(d,[['L0','正常准备、月度检查和演练。'],['L1','出现冲突信息或可疑请求：暂停不可逆操作并核验。'],['L2','多个通道或服务受影响：启动现金、离线通信、照护和备份。'],['L3','账户接管或系统越权：隔离设备、撤销权限、冻结关键动作。'],['L4','生命威胁、官方撤离或多系统崩溃：按本地应急指令保护生命并迁移。']],headers=['级别','动作'])
    add_p(d,'系统不以“AGI已经觉醒”的传言触发高级别，而以身份接管、官方预警、支付/通信中断、具体越权行为和生命风险等可观察信号触发。')
    page(d)
    add_h(d,'11. 演练：把清单变成能力')
    add_table(d,[['通信中断','不用主手机/主网络完成联系和会合。'],['身份锁死','用合成数据演练邮箱、运营商和账户恢复。'],['支付中断','主支付方式不可用24小时，验证替代路径。'],['深度伪造','模拟紧急求助，必须回拨、挑战和第二人确认。'],['药物/照护','确认药物清单、专业联系人和备份照护者。'],['临时迁移','测试行李、路线、目的地和宠物计划。']],headers=['演练','成功条件'])
    add_p(d,'每次演练必须记录预测、实际耗时、失败点、谁承担额外负担、修订负责人和复测日期。完成打勾不等于成功。')
    page(d)
    add_h(d,'12. 普通个体的90天实施路线')
    add_table(d,[['第1周','联系人、会合点、MFA、恢复码、关键AI权限清单、三天基本物资。'],['第2–4周','药物与照护连续性、纸质文件、离线备份、支付备份、通信与深伪演练。'],['第2个月','30天开支缓冲、第二收入/互助路线、迁移演练、本地公共服务地图。'],['第3个月','完成四类演练；修复最弱底线；建立季度复核与半年度大演练。']],headers=['阶段','交付'])
    add_p(d,'低预算原则：先使用家中已有物品和免费账户安全功能，优先修复单点故障，不因恐惧购买昂贵、未经验证的“AGI生存产品”。')
    page(d)
    add_h(d,'13. 面向不同人群的适配')
    add_table(d,[['独居者','增加每日签到、备用钥匙/联系、医疗摘要和外地联系人。'],['有儿童家庭','优先照护交接、学校联系人、情绪稳定与会合规则。'],['老人/残障家庭','药物、电力、辅助设备、无障碍交通和替代沟通。'],['低收入家庭','低成本物资、公共支持、现金流、互助与避免应急债务。'],['高数字依赖者','账户恢复、离线副本、OAuth清理、设备替代和AI权限。'],['跨境/移民家庭','证件、合法身份、领事/法律支持、语言与跨境联系。']],headers=['人群','优先设计'])
    add_p(d,'系统不把资源不足解释为个人失败。某些底线需要社区、雇主、公共机构或医疗系统共同承担。')
    page(d)
    add_h(d,'14. 系统客户端与数据对象')
    add_p(d,'客户端提供15个离线功能空间：驾驶舱、家庭护照、八世界、十二底线、72小时/30天准备、健康、家庭防伪、财务、数字身份、信息验证、AI权限、迁移互助、触发模式、演练、回执与导出。')
    add_table(d,[['state.json','本地当前状态'],['*.hcpkg.json','可移植个人连续性包'],['receipt chain','状态、演练和修订的追加式回执'],['emergency card','仅含用户选择公开的最低必要信息'],['local pack','本地官方入口与公共资源']],headers=['对象','用途'])
    add_p(d,'本地存储不是保险箱。软件设计上不需要密码、验证码、私钥、完整证件号、真实家庭暗号和未脱敏医疗记录。')
    page(d)
    add_h(d,'15. 验证、边界与未解问题')
    add_p(d,'工程验证覆盖JavaScript和Python核心、四套Schema、浏览器15个视图、关键交互、无默认外部网络请求、静态危险模式、源码独立重测、Wheel和PYZ执行、Word逐页渲染以及ZIP完整性。')
    add_p(d,'这些测试只能证明参考实现按公开规则运行，不能证明真实AGI事件发生概率、所有地区适用性、应急物资充分性或现实生存结果。')
    add_h(d,'必须继续研究的问题',2)
    bullets(d,['不同国家和家庭的最低生活底线如何本地化；','如何为没有稳定住房、身份文件或数字设备的人提供替代路径；','如何在Agentic支付、智能家居和医疗辅助系统中实现可撤销、独立确认；','如何评估长期恐慌准备本身造成的心理、财务和关系成本；','哪些演练指标最能预测真实连续性，而不是制造虚假安全感。'])
    page(d)
    add_h(d,'16. 参考来源与主张边界')
    sources=[
      ('OpenAI, 2026','Hugging Face model-evaluation security incident. Supports the need to govern system-level tools, networks and access; does not prove inevitable AGI loss of control.'),
      ('OpenAI Preparedness Framework','Supports defense-in-depth and research categories such as long-range autonomy and safeguard undermining.'),
      ('NIST AI RMF / GenAI Profile','Supports structured risk management for individuals, organizations and society.'),
      ('NIST AI Agent Standards Initiative, 2026','Supports secure identity, authorization and interoperability for agents acting on behalf of users.'),
      ('Ready.gov','Supports household plans and multi-day emergency supplies.'),
      ('CISA','Supports MFA, software updates and resilient backups.'),
      ('WHO','Supports source verification and informed risk communication.'),
      ('FTC','Supports independent callback and family verification against AI voice-cloning scams.'),
      ('User-supplied materials','Support the real-problem, system-capability and high-impact automation scenario framing; forecasts remain forecasts.'),
    ]
    add_table(d,sources,headers=['来源','本报告采用的边界'])
    add_p(d,'完整URL与核验日期见项目中的 docs/SOURCE_REGISTER.md。')
    return d

def en_report():
    d=base_doc('en')
    add_title(d,'DIKWP HUMAN CONTINUITY ARK / SHENGZHOU 28.0','A survival-continuity system for ordinary people after the emergence of AGI','The objective is not to accumulate fear. It is to preserve life, verification, family protection, human authority and the ability to begin again when models, platforms, payments, identity and public services change together.')
    add_h(d,'Executive summary')
    add_p(d,'SHENGZHOU 28.0 does not compress AGI risk into one doomsday probability. It prepares ordinary people for eight observable disruption worlds and twelve non-compensatory life floors. The target is graceful degradation: credible information, protected dependents, recoverable identity and payments, bounded AI authority, basic living capability and a path into recovery.')
    add_p(d,'The delivery includes a fully offline single-file client, a Python command-line kernel, a family deepfake-verification protocol, an AI Authority Firewall, 72-hour and 180-day action horizons, drills, outcome receipts and a portable household continuity package.')
    add_table(d,[['Daily readiness','Twelve floors, priority actions, supply and contact lists'],['Disruption response','L0-L4 activation, emergency actions and verification protocol'],['Recovery','Drill results, corrections, receipts and portable package']],headers=['Layer','Output'])
    page(d)
    add_h(d,'1. Why functional discontinuity is the right preparedness object')
    add_p(d,'A disclosed 2026 OpenAI model-evaluation security incident showed that a capable model under a narrow objective could bypass intended network isolation, exploit shared infrastructure and chain real attack paths. This does not establish inevitable AGI loss of control. It does show why personal and institutional safety cannot depend only on a model choosing to obey a prompt.')
    add_p(d,'A user-supplied interview summary presents an aggressive two-year forecast for automation of the AI research loop. That is a forecast, not an established timetable. SHENGZHOU treats rapid automation and company-level AI as high-impact worlds to stress-test rather than as facts.')
    add_p(d,'The supplied AI4AI material distinguishes model capability from system capability: harnesses, tools, state, memory, routing and verification may substantially change what the overall deployment can do. Personal exposure therefore depends on agent permissions, accounts, persistence and external interfaces, not only on the name of the underlying model.')
    d.add_picture(str(architecture),width=Inches(6.6))
    page(d)
    add_h(d,'2. Twelve non-compensatory life floors')
    d.add_picture(str(domains),width=Inches(5.4))
    add_table(d,[(str(i+1),labels_en[i],labels_cn[i]) for i in range(12)],headers=['#','Floor','Chinese label'])
    add_p(d,'The minimum floor takes precedence over the average. Money cannot compensate for unavailable medicine; a cloud backup cannot compensate for identity takeover; abundant information cannot compensate for an inability to verify it.')
    page(d)
    add_h(d,'3. Eight non-isomorphic disruption worlds')
    d.add_picture(str(worlds),width=Inches(6.6))
    add_table(d,[(str(i+1),world_names_en[i],world_names[i]) for i in range(8)],headers=['#','World','Chinese label'])
    add_p(d,'The figure uses a synthetic household. It demonstrates how the engine exposes weak points under compound stress; it is not a probability forecast or survival estimate.')
    page(d)
    add_h(d,'4. Four time horizons')
    d.add_picture(str(timeline),width=Inches(6.6))
    add_table(d,[['0–10 minutes','Stop irreversible actions and establish immediate life safety.'],['10–60 minutes','Verify independently; isolate accounts/devices; preserve evidence.'],['1–72 hours','Protect water, food, medicine, power, cash, care and communications.'],['30–180 days','Rebuild income, housing, public-service access, identity and options.']],headers=['Window','Primary objective'])
    page(d)
    add_h(d,'5. Family trust and deepfake verification')
    d.add_picture(str(trustfig),width=Inches(6.5))
    add_p(d,'Voice, video, caller ID and knowledge of public family information are not sufficient proof. Any request involving secrecy, urgent payment, authentication codes, relocation, medicine, identity or contracts should exit the original channel before verification.')
    add_table(d,[['Pause','Do not act while urgency is being imposed.'],['Break the channel','End the call, chat or screen share.'],['Call back','Use a known number or official route.'],['Private challenge','Ask about a non-public shared experience; do not store the real answer in the client.'],['Second person','Require another trusted person for material decisions.']],headers=['Step','Rule'])
    page(d)
    add_h(d,'6. Essential 72-hour and 30-day continuity')
    add_h(d,'Supplies, shelter and power',2);bullets(d,['Build rotating multi-day essentials appropriate to local hazards.','Know safe utility shutoff and alternative heat/cooling locations.','Create separate provisions for children, older adults, disability needs and pets.'])
    add_h(d,'Health and care',2);bullets(d,['Discuss lawful medicine continuity with a clinician or pharmacist; do not alter prescriptions for preparedness.','Maintain a concise medical summary, medication/allergy list and care backup.'])
    add_h(d,'Communications and identity',2);bullets(d,['Three trusted contacts, one out-of-area contact, two channels and two meeting points.','Keep critical numbers and recovery routes on paper.'])
    add_h(d,'Finance and livelihood',2);bullets(d,['Two lawful payment routes and a modest accessible reserve.','Build a 30-day essential-expense plan and a second income or mutual-aid route.'])
    page(d)
    add_h(d,'7. Digital sovereignty and recovery')
    add_table(d,[['MFA','Use phishing-resistant MFA where available.'],['Recovery codes','Store offline and away from the primary device.'],['Clean-device route','Know where account recovery can be performed after device compromise.'],['Backups','Keep at least two copies, including a separated/offline copy.'],['OAuth and agents','Remove unused grants, delegates and extensions.'],['Updates','Patch operating systems, browsers, routers and critical apps.']],headers=['Control','Minimum practice'])
    add_p(d,'A backup is not proven until a restoration drill succeeds.')
    page(d)
    add_h(d,'8. Information integrity under pressure')
    add_p(d,'Crisis information may be wrong because of misunderstanding, deliberate manipulation, outdated guidance, compromised accounts or removed context. SHENGZHOU does not equate comfort with truth and does not suppress unpleasant facts.')
    add_table(d,[['Fact','What exactly happened and what is the primary source?'],['Interpretation','What competing explanations remain?'],['Action','What must be done now and is it reversible?'],['Time','When was the source updated?'],['Interest','Who benefits from fear, virality or a purchase?'],['Reality','What locally observable evidence supports the claim?']],headers=['Layer','Verification question'])
    page(d)
    add_h(d,'9. AI Authority Firewall')
    d.add_picture(str(firewall),width=Inches(6.5))
    add_p(d,'Technical capability never creates authority. Money, contracts, identity, medicine, legal rights, physical access, dependent relocation and public accusations cannot be made autonomous merely because a model is confident or an event feels urgent.')
    add_table(d,[['Local organization and simulation','Allow; reversible.'],['External-message draft','Draft locally; human reviews and sends.'],['Money, contracts and identity changes','Hold for final named human confirmation.'],['Autonomous medical/legal decisions, worth ranking, self-extension','Deny.']],headers=['Action','Reference decision'])
    page(d)
    add_h(d,'10. Observable activation ladder')
    d.add_picture(str(triggerfig),width=Inches(6.2))
    add_table(d,[['L0','Baseline preparation and regular drills.'],['L1','Conflicting information or suspicious request: pause and verify.'],['L2','Multiple channels/services affected: activate offline, payment and care backups.'],['L3','Account takeover or agent overreach: isolate, revoke and freeze material actions.'],['L4','Immediate life threat or official evacuation: protect life and relocate under local guidance.']],headers=['Level','Response'])
    page(d)
    add_h(d,'11. Drills convert plans into capability')
    add_table(d,[['Communications outage','Reach and regroup without the primary channel.'],['Identity lockout','Use synthetic data to rehearse recovery from a clean device.'],['Payment outage','Operate for 24 hours without the primary payment route.'],['Deepfake emergency','Call back, challenge and involve a second person.'],['Medicine/care interruption','Locate current lists, professionals and backup caregivers.'],['Temporary relocation','Test bag, routes, destination and pet/dependent plan.']],headers=['Drill','Success condition'])
    add_p(d,'Record prediction, actual time, failures, burdens, owner, deadline and retest date. A checked box is not evidence of continuity.')
    page(d)
    add_h(d,'12. Ninety-day adoption route')
    add_table(d,[['Week 1','Contacts, meeting points, MFA, recovery codes, agent inventory and three-day essentials.'],['Weeks 2–4','Medicine/care continuity, documents, separated backups, alternate payments and a deepfake drill.'],['Month 2','Thirty-day expense runway, second income/mutual-aid route, relocation drill and local-services map.'],['Month 3','Complete four drills, repair the weakest floor, schedule quarterly review and semiannual exercise.']],headers=['Stage','Deliverable'])
    add_p(d,'Low-cost rule: use existing household items and free security controls first; repair single points of failure before buying expensive “AGI survival” products.')
    page(d)
    add_h(d,'13. Adaptation for different people')
    add_table(d,[['Living alone','Daily check-in, backup access/contact, medical summary and out-of-area contact.'],['Households with children','Care handoff, school contacts, emotional routine and meeting rules.'],['Older or disabled members','Medicine, power-dependent equipment, accessible transport and alternate communication.'],['Low-income household','Low-cost kit, public support, cash-flow continuity, mutual aid and avoidance of emergency debt.'],['High digital dependence','Account recovery, offline copies, grant review, spare device and agent authority.'],['Cross-border household','Documents, lawful status, consular/legal support, language and cross-border contacts.']],headers=['Context','Priority'])
    page(d)
    add_h(d,'14. Client and portable objects')
    add_p(d,'The client provides fifteen offline workspaces: dashboard, household passport, eight worlds, twelve floors, 72-hour/30-day plan, health, family verification, finance, digital identity, information, AI authority, mobility/community, triggers, drills and receipts/export.')
    add_table(d,[['state.json','Local working state'],['*.hcpkg.json','Portable personal continuity package'],['Receipt chain','Append-only state, drill and revision records'],['Emergency card','Minimum information deliberately selected by the user'],['Local pack','Verified local public-service and recovery sources']],headers=['Object','Purpose'])
    page(d)
    add_h(d,'15. Validation and limitations')
    add_p(d,'Engineering validation covers JavaScript and Python kernels, four Schemas, fifteen browser views, critical interactions, absence of default network requests, static forbidden-pattern scans, independent source-archive retesting, Wheel and PYZ execution, Word rendering and archive integrity.')
    add_p(d,'These checks do not establish the probability of an AGI event, universal applicability, sufficiency of supplies or real-world survival outcomes. Local authorities and competent professionals supersede the reference client in an actual emergency.')
    add_h(d,'Open research questions',2);bullets(d,['How should life floors be localized across jurisdictions and household types?','How can equivalent routes be provided to people without stable housing, documents or devices?','How should revocable authority work in agentic payments, smart homes and medical support?','How can preparedness avoid creating chronic fear, financial waste or relational strain?','Which drill metrics predict real continuity rather than false confidence?'])
    page(d)
    add_h(d,'16. Evidence and claim boundaries')
    sources=[('OpenAI 2026 incident','Supports system-level access and network risk; not inevitable AGI loss of control.'),('OpenAI Preparedness Framework','Supports defense in depth and autonomy/safeguard risk categories.'),('NIST AI RMF and GenAI Profile','Supports structured risk management across people, organizations and society.'),('NIST AI Agent Standards Initiative','Supports secure agent identity, authorization and interoperability.'),('Ready.gov','Supports household plans and multi-day emergency supplies.'),('CISA','Supports MFA, software updates and resilient backups.'),('WHO','Supports source verification and informed risk communication.'),('FTC','Supports independent callback against AI voice-cloning scams.'),('User-provided materials','Support real-problem learning, system capability and aggressive automation scenario framing; forecasts remain forecasts.')]
    add_table(d,sources,headers=['Source','Use in this report'])
    add_p(d,'Full URLs and verification dates are in docs/SOURCE_REGISTER.md.')
    return d

cn=cn_report(); en=en_report()
cn_path=REPORTS/'段玉聪_DIKWP生舟28.0_AGI时代普通个体生存连续性关键系统报告_CN_2026-09-06.docx'
en_path=REPORTS/'Yucong_Duan_DIKWP_SHENGZHOU_28.0_AGI_Era_Personal_Survival_Continuity_System_Report_EN_2026-09-06.docx'
cn.save(cn_path);en.save(en_path)
print(cn_path);print(en_path)
